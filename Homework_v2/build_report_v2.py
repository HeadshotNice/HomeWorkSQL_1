from pathlib import Path
import shutil

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


BASE = Path(__file__).resolve().parent
FIG_DIR = BASE / "figures"
FIG_DIR.mkdir(exist_ok=True)


def font(path, size):
    try:
        return ImageFont.truetype(path, size)
    except Exception:
        return ImageFont.load_default()


def make_er_diagram():
    img_path = FIG_DIR / "er_diagram_v2.png"
    image = Image.new("RGB", (1800, 1050), "white")
    draw = ImageDraw.Draw(image)
    font_title = font(r"C:\Windows\Fonts\arialbd.ttf", 34)
    font_box = font(r"C:\Windows\Fonts\arialbd.ttf", 24)
    font_text = font(r"C:\Windows\Fonts\arial.ttf", 20)
    font_small = font(r"C:\Windows\Fonts\arial.ttf", 17)
    blue = (43, 88, 134)
    gray = (90, 90, 90)
    green = (42, 117, 92)

    def box(x, y, w, h, title, lines):
        draw.rounded_rectangle([x, y, x + w, y + h], radius=16, outline=blue, width=3, fill=(232, 241, 250))
        draw.rectangle([x, y, x + w, y + 44], fill=blue)
        draw.text((x + 14, y + 10), title, fill="white", font=font_box)
        yy = y + 58
        for line in lines:
            draw.text((x + 14, yy), line, fill=(30, 30, 30), font=font_small)
            yy += 25
        return (x, y, x + w, y + h)

    def center(rect):
        x1, y1, x2, y2 = rect
        return ((x1 + x2) // 2, (y1 + y2) // 2)

    def line(a, b):
        ax, ay = center(a)
        bx, by = center(b)
        draw.line([ax, ay, bx, by], fill=gray, width=3)

    draw.text((45, 30), "ER-диаграмма предметной области: прокат автомобилей", fill=blue, font=font_title)
    positions = {
        "BRANCH": (70, 140, 360, 180),
        "EMPLOYEE": (560, 125, 360, 205),
        "CUSTOMER": (70, 430, 360, 205),
        "CAR": (560, 430, 360, 235),
        "RENTAL": (1070, 320, 380, 260),
        "PAYMENT": (1070, 705, 380, 175),
        "MAINT": (560, 760, 360, 185),
    }
    rects = {key: (x, y, x + w, y + h) for key, (x, y, w, h) in positions.items()}
    for a, b in [
        ("BRANCH", "EMPLOYEE"),
        ("BRANCH", "CAR"),
        ("CUSTOMER", "RENTAL"),
        ("CAR", "RENTAL"),
        ("EMPLOYEE", "RENTAL"),
        ("RENTAL", "PAYMENT"),
        ("CAR", "MAINT"),
        ("EMPLOYEE", "MAINT"),
    ]:
        line(rects[a], rects[b])
    rects = {
        "BRANCH": box(70, 140, 360, 180, "BRANCH / Филиал", ["candidate key: branch_name", "city, address, phone"]),
        "EMPLOYEE": box(560, 125, 360, 205, "EMPLOYEE / Сотрудник", ["candidate key: personnel_no", "role, branch, full_name", "hired_at, fired_at"]),
        "CUSTOMER": box(70, 430, 360, 205, "CUSTOMER / Клиент", ["candidate key: license_no", "full_name, phone, email", "license dates"]),
        "CAR": box(560, 430, 360, 235, "CAR / Автомобиль", ["candidate keys: plate_no, VIN", "model characteristics*", "branch, color, odometer", "active flag"]),
        "RENTAL": box(1070, 320, 380, 260, "RENTAL / Договор", ["candidate key: contract_no", "customer, car, branches", "agent, status, period", "daily_rate, deposit"]),
        "PAYMENT": box(1070, 705, 380, 175, "PAYMENT / Платеж", ["candidate key: payment_doc_no", "rental, method, paid_at", "amount, refund flag"]),
        "MAINT": box(560, 760, 360, 185, "MAINTENANCE / ТО", ["candidate key: service_act_no", "car, mechanic, period", "description, cost"]),
    }
    draw.rectangle([55, 948, 1540, 1020], fill="white")
    draw.text((70, 955), "Все связи на диаграмме имеют тип 1:N; обязательность и внешние ключи указаны в финальной схеме отношений.", fill=green, font=font_text)
    draw.text(
        (70, 995),
        "* Характеристики модели автомобиля на логическом этапе вынесены в отношение car_model.",
        fill=(40, 40, 40),
        font=font_text,
    )
    image.save(img_path)
    return img_path


RELATIONS = {
    "car_category": [
        ("category_code", "CHAR(4)", "PK, NOT NULL", "Код категории автомобиля"),
        ("category_name", "VARCHAR(50)", "NOT NULL, UNIQUE", "Название категории"),
    ],
    "fuel_type": [
        ("fuel_code", "CHAR(2)", "PK, NOT NULL", "Код типа топлива"),
        ("fuel_name", "VARCHAR(30)", "NOT NULL, UNIQUE", "Название типа топлива"),
    ],
    "transmission_type": [
        ("transmission_code", "CHAR(1)", "PK, NOT NULL", "Код типа КПП"),
        ("transmission_name", "VARCHAR(20)", "NOT NULL, UNIQUE", "Название типа КПП"),
    ],
    "payment_method": [
        ("method_code", "CHAR(4)", "PK, NOT NULL", "Код способа оплаты"),
        ("method_name", "VARCHAR(30)", "NOT NULL, UNIQUE", "Название способа оплаты"),
    ],
    "employee_role": [
        ("role_code", "CHAR(4)", "PK, NOT NULL", "Код роли сотрудника"),
        ("role_name", "VARCHAR(50)", "NOT NULL, UNIQUE", "Название роли"),
    ],
    "rental_status": [
        ("status_code", "CHAR(3)", "PK, NOT NULL", "Код статуса договора"),
        ("status_name", "VARCHAR(30)", "NOT NULL, UNIQUE", "Название статуса"),
    ],
    "branch": [
        ("branch_id", "BIGSERIAL", "PK", "Суррогатный ключ филиала"),
        ("branch_name", "VARCHAR(100)", "NOT NULL, UNIQUE", "Название филиала"),
        ("city", "VARCHAR(60)", "NOT NULL", "Город"),
        ("address", "VARCHAR(200)", "NOT NULL", "Адрес"),
        ("phone", "VARCHAR(20)", "NULL", "Телефон"),
    ],
    "employee": [
        ("employee_id", "BIGSERIAL", "PK", "Суррогатный ключ сотрудника"),
        ("branch_id", "BIGINT", "FK -> branch, NOT NULL", "Филиал сотрудника"),
        ("role_code", "CHAR(4)", "FK -> employee_role, NOT NULL", "Роль сотрудника"),
        ("full_name", "VARCHAR(120)", "NOT NULL", "ФИО"),
        ("phone", "VARCHAR(20)", "NULL", "Телефон"),
        ("hired_at", "DATE", "NOT NULL", "Дата приема"),
        ("fired_at", "DATE", "NULL, CHECK", "Дата увольнения не раньше приема"),
    ],
    "customer": [
        ("customer_id", "BIGSERIAL", "PK", "Суррогатный ключ клиента"),
        ("full_name", "VARCHAR(120)", "NOT NULL", "ФИО"),
        ("phone", "VARCHAR(20)", "NOT NULL", "Телефон"),
        ("email", "VARCHAR(254)", "NULL, CHECK", "E-mail содержит @"),
        ("birth_date", "DATE", "NULL", "Дата рождения"),
        ("driver_license_no", "VARCHAR(32)", "NOT NULL, UNIQUE", "Номер водительского удостоверения"),
        ("driver_license_issued", "DATE", "NULL", "Дата выдачи прав"),
        ("driver_license_expires", "DATE", "NULL, CHECK", "Дата окончания прав"),
    ],
    "car_model": [
        ("model_id", "BIGSERIAL", "PK", "Суррогатный ключ модели"),
        ("make", "VARCHAR(40)", "NOT NULL", "Марка"),
        ("model", "VARCHAR(60)", "NOT NULL", "Модель"),
        ("year", "SMALLINT", "NOT NULL, CHECK", "Год выпуска модели"),
        ("category_code", "CHAR(4)", "FK -> car_category, NOT NULL", "Категория"),
        ("fuel_code", "CHAR(2)", "FK -> fuel_type, NOT NULL", "Тип топлива"),
        ("transmission_code", "CHAR(1)", "FK -> transmission_type, NOT NULL", "Тип КПП"),
        ("seats", "SMALLINT", "NOT NULL, CHECK", "Количество мест"),
        ("UNIQUE", "(make, model, year, category_code, fuel_code, transmission_code)", "UNIQUE", "Естественная уникальность комплектации"),
    ],
    "car": [
        ("car_id", "BIGSERIAL", "PK", "Суррогатный ключ автомобиля"),
        ("model_id", "BIGINT", "FK -> car_model, NOT NULL", "Модель"),
        ("current_branch_id", "BIGINT", "FK -> branch, NOT NULL", "Текущий филиал"),
        ("plate_no", "VARCHAR(15)", "NOT NULL, UNIQUE", "Госномер"),
        ("vin", "CHAR(17)", "NOT NULL, UNIQUE, CHECK", "VIN"),
        ("color", "VARCHAR(30)", "NULL", "Цвет"),
        ("odometer_km", "INTEGER", "NOT NULL, CHECK", "Пробег неотрицательный"),
        ("active", "BOOLEAN", "NOT NULL", "Можно ли сдавать автомобиль"),
    ],
    "rental": [
        ("rental_id", "BIGSERIAL", "PK", "Суррогатный ключ договора"),
        ("customer_id", "BIGINT", "FK -> customer, NOT NULL", "Клиент"),
        ("car_id", "BIGINT", "FK -> car, NOT NULL", "Автомобиль"),
        ("pickup_branch_id", "BIGINT", "FK -> branch, NOT NULL", "Филиал выдачи"),
        ("dropoff_branch_id", "BIGINT", "FK -> branch, NOT NULL", "Филиал возврата"),
        ("opened_by_employee_id", "BIGINT", "FK -> employee, NOT NULL", "Агент, открывший договор"),
        ("status_code", "CHAR(3)", "FK -> rental_status, NOT NULL", "Статус договора"),
        ("planned_period", "TSTZRANGE", "NOT NULL, CHECK, EXCLUDE", "Плановый период аренды"),
        ("actual_period", "TSTZRANGE", "NULL, CHECK", "Фактический период"),
        ("daily_rate", "NUMERIC(10,2)", "NOT NULL, CHECK", "Суточный тариф > 0"),
        ("deposit_amount", "NUMERIC(10,2)", "NOT NULL, CHECK", "Залог >= 0"),
        ("notes", "VARCHAR(500)", "NULL", "Примечания"),
    ],
    "payment": [
        ("payment_id", "BIGSERIAL", "PK", "Суррогатный ключ платежа"),
        ("rental_id", "BIGINT", "FK -> rental, NOT NULL", "Договор"),
        ("method_code", "CHAR(4)", "FK -> payment_method, NOT NULL", "Способ оплаты"),
        ("paid_at", "TIMESTAMPTZ", "NOT NULL", "Дата и время платежа"),
        ("amount", "NUMERIC(12,2)", "NOT NULL, CHECK", "Сумма > 0"),
        ("is_refund", "BOOLEAN", "NOT NULL", "Признак возврата"),
    ],
    "maintenance": [
        ("maintenance_id", "BIGSERIAL", "PK", "Суррогатный ключ обслуживания"),
        ("car_id", "BIGINT", "FK -> car, NOT NULL", "Автомобиль"),
        ("performed_by_employee_id", "BIGINT", "FK -> employee, NULL", "Механик"),
        ("service_period", "TSTZRANGE", "NOT NULL, CHECK, EXCLUDE", "Период обслуживания"),
        ("odometer_km", "INTEGER", "NULL, CHECK", "Пробег на момент ТО"),
        ("description", "VARCHAR(400)", "NOT NULL", "Описание работ"),
        ("cost_amount", "NUMERIC(12,2)", "NULL, CHECK", "Стоимость >= 0"),
    ],
}

VIEWS = [
    ("v_cars", "Карточка автомобиля с моделью, категорией, филиалом и признаками доступности."),
    ("v_active_rentals", "Активные и новые договоры для агента и менеджера."),
    ("v_overdue_rentals", "Договоры, по которым плановый срок уже прошел, но статус еще активный/новый."),
    ("v_available_cars_now", "Автомобили, доступные в текущий момент."),
    ("v_rental_payment_net", "Сальдо платежей по договору с учетом возвратов."),
    ("v_rentals_pricing", "Расчет плановых дней аренды, стоимости и оплаченной суммы."),
    ("v_cars_in_maintenance", "Автомобили, находящиеся на обслуживании сейчас."),
    ("v_revenue_by_month", "Выручка по месяцам по данным платежей."),
    ("v_customer_rental_history", "История договоров клиента."),
    ("v_car_utilization", "Плановая загрузка автомобиля в днях."),
]

CONSTRAINTS = [
    ("Непересечение аренд одного автомобиля", "EXCLUDE USING gist по car_id и planned_period", "Отмененные договоры CNL не блокируют период."),
    ("Непересечение обслуживаний одного автомобиля", "EXCLUDE USING gist по car_id и service_period", "Не позволяет завести два ТО на один период."),
    ("Аренда не пересекается с обслуживанием", "Два триггера + pg_advisory_xact_lock", "Проверка сериализуется по car_id."),
    ("Неактивный автомобиль нельзя сдавать", "trg_car_must_be_active", "Проверка при INSERT и UPDATE для NEW/ACT."),
    ("Права клиента действуют на период аренды", "trg_customer_license_valid_for_rental", "Проверка относительно planned_period, без CHECK от current_date."),
    ("Договор открывает агент или менеджер", "trg_rental_opened_by_agent_or_manager", "ACCT/MECH не могут быть opened_by_employee_id."),
    ("Обслуживание выполняет механик", "trg_maintenance_performed_by_mechanic", "Если механик указан, роль должна быть MECH."),
    ("Положительные суммы и корректные даты", "CHECK constraints", "Простые ОЦ реализованы без лишних триггеров."),
]

ROLES = [
    ("cr_manager", "Менеджер сети", "Полный доступ к объектам схемы", "Контроль справочников, отчетов, структуры и качества данных."),
    ("cr_agent", "Агент проката", "SELECT справочников и авто; INSERT/UPDATE customer, rental; нужные sequence", "Регистрация клиентов, создание и изменение договоров."),
    ("cr_accountant", "Бухгалтер", "SELECT договоров/клиентов; INSERT/UPDATE payment; payment sequence", "Учет платежей и возвратов, финансовые отчеты."),
    ("cr_mechanic", "Механик/сервис", "SELECT автомобилей; INSERT/UPDATE maintenance; maintenance sequence", "Регистрация работ по обслуживанию и просмотр ТО."),
]

KPI_ROWS = [
    ("Формат сдачи", 10, 8, "Создан DOCX v2; осталось заменить плейсхолдеры фамилий/группы."),
    ("Инфологическое проектирование", 15, 14, "Есть описание ПрО, документы, пользователи, ER-диаграмма и сущности."),
    ("Логическая схема и нормализация", 25, 23, "Добавлены отношения с типами, ключами, ФЗ и нормализация до 4НФ."),
    ("Ограничения целостности", 15, 14, "CHECK, FK, EXCLUDE и триггеры; конкурентные проверки усилены advisory lock."),
    ("Физическая реализация SQL", 15, 13, "Скрипт расширен; нужен финальный прогон в живом PostgreSQL."),
    ("Представления и индексы", 10, 10, "10 представлений и индексы по всем FK + составной индекс."),
    ("Права и резервное копирование", 10, 9, "Права стали точечнее, стратегия РК с числом копий описана."),
]


def write_markdown(img_path):
    lines = [
        "# Домашнее задание по курсу «Базы данных»",
        "## Тема: БД фирмы по прокату автомобилей",
        "",
        "**Исполнители:** <Фамилия Имя>, <Фамилия Имя>  ",
        "**Группа:** <номер группы>  ",
        "**Версия:** 2  ",
        "**СУБД:** PostgreSQL",
        "",
        "> Перед отправкой заменить плейсхолдеры в титуле и имени файла на реальные фамилии, группу и номер версии.",
        "",
        "## 2.1. Инфологическое проектирование",
        "### 2.1.1. Анализ предметной области",
        "Организация занимается прокатом автомобилей через сеть филиалов. Клиент выбирает автомобиль, заключает договор аренды на плановый период, получает автомобиль в филиале выдачи и возвращает его в том же или другом филиале. Сотрудники оформляют договоры, бухгалтерия регистрирует оплаты и возвраты, сервисная служба фиксирует обслуживание автомобилей.",
        "",
        "**Основные документы предметной области:** карточка филиала, карточка сотрудника, карточка клиента, карточка автомобиля, договор аренды, платежный документ, акт обслуживания автомобиля.",
        "",
        "**Бизнес-правила:**",
        "- Один автомобиль не может быть сдан двум клиентам на пересекающиеся периоды.",
        "- Автомобиль на обслуживании недоступен для аренды на пересекающийся период.",
        "- Договор имеет статус: новый, активный, закрыт, отменен.",
        "- Оплаты относятся к конкретному договору; платеж может быть возвратом.",
        "- Договор может открыть только агент проката или менеджер.",
        "- Обслуживание автомобиля выполняет сотрудник с ролью механика.",
        "",
        "### 2.1.2. Пользователи и информационные задачи",
        "| Роль | Пользователь | Права и задачи | Обоснование |",
        "|---|---|---|---|",
    ]
    lines += [f"| {' | '.join(row)} |" for row in ROLES]
    lines += [
        "",
        "### 2.1.3. ER-диаграмма",
        f"![ER-диаграмма](figures/{img_path.name})",
        "",
        "## 2.2. Требования к операционной обстановке",
        "База рассчитана на небольшую сеть проката: несколько филиалов, десятки сотрудников, сотни или тысячи клиентов и автомобилей. Требуется многопользовательский доступ, транзакционность и надежное восстановление после сбоя.",
        "",
        "## 2.3. Выбор СУБД",
        "Выбрана PostgreSQL: она поддерживает FK, транзакции, роли, представления, PL/pgSQL, GiST-индексы, exclusion constraints и диапазонные типы `tstzrange`.",
        "",
        "## 2.4. Логическое проектирование реляционной БД",
        "### 2.4.1. Преобразование ER-диаграммы в схему БД",
        "Все связи 1:М реализованы внешними ключами на стороне «много». Связей М:М в итоговой схеме нет. Справочные атрибуты вынесены в отдельные отношения. Суррогатные ключи *_id введены на логическом этапе для компактных FK; на ER-диаграмме показаны потенциальные предметные ключи.",
        "",
        "### 2.4.2. Финальная схема отношений",
    ]
    for name, rows in RELATIONS.items():
        lines += [f"#### `{name}`", "| Поле | Тип | Ключ/обязательность | Назначение и ограничения |", "|---|---|---|---|"]
        lines += [f"| {' | '.join(row)} |" for row in rows]
        lines.append("")
    lines += [
        "### 2.4.3. Нормализация до 4НФ",
        "`RENTAL_RAW(Договор, КлиентФИО, КлиентТелефон, КлиентПрава, АвтоНомер, VIN, Марка, Модель, Категория, Топливо, КПП, ФилиалВыдачи, ФилиалВозврата, Агент, Статус, Период, Тариф, ПлатежСумма, ПлатежМетод, ТОПериод, ТООписание, ...)`",
        "",
        "Функциональные зависимости: `driver_license_no -> данные клиента`, `plate_no -> данные автомобиля`, `model_id -> характеристики модели`, `rental_id -> данные договора`, `payment_id -> платеж`, `maintenance_id -> обслуживание`.",
        "",
        "1НФ достигается атомарностью атрибутов и выносом повторяющихся платежей/ТО. 2НФ выполняется, так как неключевые атрибуты зависят от полного ключа. 3НФ/BCNF достигается выносом клиента, автомобиля, модели, справочников, платежей и обслуживания. 4НФ выполняется, потому что независимые многозначные факты не хранятся в одном отношении.",
        "",
        "### 2.4.4. Дополнительные ограничения целостности",
        "| Ограничение | Реализация | Комментарий |",
        "|---|---|---|",
    ]
    lines += [f"| {' | '.join(row)} |" for row in CONSTRAINTS]
    lines += [
        "",
        "### 2.4.5. Права доступа",
        "| Роль | Пользователь | Права | Обоснование |",
        "|---|---|---|---|",
    ]
    lines += [f"| {' | '.join(row)} |" for row in ROLES]
    lines += [
        "",
        "## 2.5. Реализация проекта БД",
        "### 2.5.1. Таблицы",
        "Таблицы, ключи, CHECK-ограничения и справочники реализованы в `car_rental.sql`. Заполнение данными выполнено только для справочников.",
        "",
        "### 2.5.2. Представления",
        "| Представление | Назначение |",
        "|---|---|",
    ]
    lines += [f"| {' | '.join(row)} |" for row in VIEWS]
    lines += [
        "",
        "### 2.5.3. Назначение прав доступа",
        "Права назначаются через `CREATE ROLE` и `GRANT`; доступ к sequence у обычных ролей точечный.",
        "",
        "### 2.5.4. Триггеры, процедуры и функции",
        "Реализованы функции: `trg_rental_period_not_in_maintenance`, `trg_maintenance_not_overlap_rentals`, `trg_rental_set_actual_period`, `trg_car_must_be_active`, `trg_customer_license_valid_for_rental`, `trg_rental_opened_by_agent_or_manager`, `trg_maintenance_performed_by_mechanic`.",
        "",
        "### 2.5.5. Индексы",
        "Созданы индексы по всем внешним ключам, по дате платежа и составной индекс `rental(car_id, status_code)`. Индексы PK/UNIQUE не дублируются.",
        "",
        "### 2.5.6. Резервное копирование",
        "Полная резервная копия выполняется ежедневно ночью. При активной эксплуатации включается WAL-архивация. Хранить 14 ежедневных и 3 еженедельные полные копии, всего 17 точек восстановления.",
        "",
        "## Приложение А. Проверка работоспособности",
        "Выполнить `car_rental.sql`, затем `smoke_tests.sql`. Ожидаемые ошибочные вставки должны быть отклонены.",
        "",
        "```bash",
        "psql -d <database> -f car_rental.sql",
        "psql -d <database> -f smoke_tests.sql",
        "```",
        "",
        "## KPI готовности v2",
        "| Критерий | Вес | Балл v2 | Комментарий |",
        "|---|---:|---:|---|",
    ]
    lines += [f"| {a} | {b} | {c} | {d} |" for a, b, c, d in KPI_ROWS]
    lines.append(f"| **Итого** | **100** | **{sum(r[2] for r in KPI_ROWS)}** | Финальный прогон в PostgreSQL нужен для закрытия технического риска. |")
    (BASE / "report.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def write_kpi():
    lines = [
        "# KPI проверки v2",
        "",
        "| Критерий | Вес | Балл v2 | Комментарий |",
        "|---|---:|---:|---|",
    ]
    lines += [f"| {a} | {b} | {c} | {d} |" for a, b, c, d in KPI_ROWS]
    lines.append(f"| **Итого** | **100** | **{sum(r[2] for r in KPI_ROWS)}** | Цель после финального ручного запуска PostgreSQL: 95+. |")
    lines += [
        "",
        "## Осталось перед отправкой",
        "",
        "- Заменить плейсхолдеры фамилий и группы в DOCX и имени файла.",
        "- Запустить `car_rental.sql` и `smoke_tests.sql` в PostgreSQL.",
        "- Прикрепить преподавателю DOCX и при необходимости SQL-файл с тем же шаблоном имени.",
    ]
    (BASE / "KPI_v2.md").write_text("\n".join(lines) + "\n", encoding="utf-8")


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, size=8.5):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, "D9EAF7")
        set_cell_text(cell, header, bold=True, size=8.8)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value, size=8.2)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Cm(width)
    doc.add_paragraph("")


def add_para(doc, text="", bold=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.name = "Times New Roman"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    run.font.size = Pt(12)
    return paragraph


def write_docx(img_path):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)
    styles = doc.styles
    styles["Normal"].font.name = "Times New Roman"
    styles["Normal"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    styles["Normal"].font.size = Pt(12)
    for style_name in ["Heading 1", "Heading 2", "Heading 3"]:
        styles[style_name].font.name = "Times New Roman"
        styles[style_name]._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        styles[style_name].font.color.rgb = RGBColor(31, 78, 121)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("Домашнее задание по курсу «Базы данных»")
    r.bold = True
    r.font.size = Pt(18)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("БД фирмы по прокату автомобилей")
    r.bold = True
    r.font.size = Pt(16)
    r.font.name = "Times New Roman"
    r._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    for text in ["Исполнители: <Фамилия Имя>, <Фамилия Имя>", "Группа: <номер группы>", "Версия: 2", "СУБД: PostgreSQL"]:
        p = doc.add_paragraph(text)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = add_para(doc, "Примечание: перед отправкой заменить плейсхолдеры в титуле и имени файла на реальные фамилии, группу и номер версии.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_page_break()

    doc.add_heading("2.1. Инфологическое проектирование", level=1)
    doc.add_heading("2.1.1. Анализ предметной области", level=2)
    add_para(doc, "Организация занимается прокатом автомобилей через сеть филиалов. Клиент выбирает автомобиль, заключает договор аренды на плановый период, получает автомобиль в филиале выдачи и возвращает его в том же или другом филиале. Сотрудники оформляют договоры, бухгалтерия регистрирует оплаты и возвраты, сервисная служба фиксирует обслуживание автомобилей.")
    add_para(doc, "Основные документы предметной области: карточка филиала, карточка сотрудника, карточка клиента, карточка автомобиля, договор аренды, платежный документ, акт обслуживания автомобиля.", bold=True)
    add_table(
        doc,
        ["Бизнес-правило", "Реализация в проекте"],
        [
            ("Один автомобиль не может быть сдан двум клиентам на пересекающиеся периоды.", "EXCLUDE constraint rental_no_overlap_per_car"),
            ("Автомобиль на обслуживании недоступен для аренды.", "Триггеры аренды/обслуживания + advisory lock"),
            ("Договор имеет ограниченный набор статусов.", "Справочник rental_status"),
            ("Платеж относится к договору и может быть возвратом.", "payment.rental_id, payment.is_refund"),
            ("Договор открывает агент или менеджер.", "trg_rental_opened_by_agent_or_manager"),
            ("Обслуживание выполняет механик.", "trg_maintenance_performed_by_mechanic"),
        ],
        widths=[8, 8],
    )
    doc.add_heading("2.1.2. Пользователи и информационные задачи", level=2)
    add_table(doc, ["Роль", "Пользователь", "Права и задачи", "Обоснование"], ROLES, widths=[3, 3.5, 6, 5])
    doc.add_heading("2.1.3. ER-диаграмма", level=2)
    add_para(doc, "ER-диаграмма содержит 7 основных сущностей. Справочники и car_model появляются на логическом этапе как результат уточнения доменов и нормализации.")
    doc.add_picture(str(img_path), width=Inches(6.6))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    for title, body in [
        ("2.2. Требования к операционной обстановке", "База рассчитана на небольшую сеть проката: несколько филиалов, десятки сотрудников, сотни или тысячи клиентов и автомобилей. Требуется многопользовательский доступ, транзакционность и надежное восстановление после сбоя. Для учебной реализации достаточно локального PostgreSQL; для эксплуатации нужен сервер PostgreSQL с регулярным резервным копированием и доступом по сети."),
        ("2.3. Выбор СУБД", "Выбрана PostgreSQL, потому что она поддерживает внешние ключи, транзакции, роли, представления, PL/pgSQL, GiST-индексы и exclusion constraints. Для данной ПрО особенно полезны диапазонные типы tstzrange, позволяющие надежно проверять пересечение периодов аренды и обслуживания."),
    ]:
        doc.add_heading(title, level=1)
        add_para(doc, body)

    doc.add_heading("2.4. Логическое проектирование реляционной БД", level=1)
    doc.add_heading("2.4.1. Преобразование ER-диаграммы в схему БД", level=2)
    add_para(doc, "Все связи 1:М реализованы внешними ключами на стороне «много». Связей М:М в итоговой схеме нет. Многозначные и справочные атрибуты вынесены в отдельные отношения или справочники. Суррогатные ключи *_id введены на логическом этапе для компактных FK; на ER-диаграмме показаны потенциальные предметные ключи.")
    doc.add_heading("2.4.2. Финальная схема отношений", level=2)
    for name, rows in RELATIONS.items():
        doc.add_heading(name, level=3)
        add_table(doc, ["Поле", "Тип", "Ключ/обязательность", "Назначение и ограничения"], rows, widths=[3.2, 3.2, 4.3, 6.0])

    doc.add_heading("2.4.3. Нормализация до 4НФ", level=2)
    add_para(doc, "Исходное ненормализованное отношение: RENTAL_RAW(Договор, КлиентФИО, КлиентТелефон, КлиентПрава, АвтоНомер, VIN, Марка, Модель, Категория, Топливо, КПП, ФилиалВыдачи, ФилиалВозврата, Агент, Статус, Период, Тариф, ПлатежСумма, ПлатежМетод, ТОПериод, ТООписание, ...).")
    add_table(
        doc,
        ["Функциональная зависимость", "Результат декомпозиции"],
        [
            ("driver_license_no -> данные клиента", "customer"),
            ("plate_no -> данные конкретного автомобиля", "car"),
            ("model_id -> характеристики модели", "car_model + справочники"),
            ("rental_id -> клиент, авто, филиалы, сотрудник, статус, период, тариф", "rental"),
            ("payment_id -> договор, способ оплаты, дата, сумма, возврат", "payment"),
            ("maintenance_id -> авто, механик, период, описание, стоимость", "maintenance"),
        ],
        widths=[7, 7],
    )
    add_para(doc, "1НФ достигается атомарностью атрибутов и выносом повторяющихся платежей/ТО в отдельные отношения. 2НФ выполняется, так как неключевые атрибуты зависят от полного ключа соответствующего отношения. 3НФ/BCNF достигается выносом клиента, автомобиля, модели, справочников, платежей и обслуживания. 4НФ выполняется, потому что независимые многозначные факты договора и автомобиля не хранятся в одном отношении.")

    doc.add_heading("2.4.4. Дополнительные ограничения целостности", level=2)
    add_table(doc, ["Ограничение", "Реализация", "Комментарий"], CONSTRAINTS, widths=[5, 6, 5])
    doc.add_heading("2.4.5. Описание групп пользователей и прав доступа", level=2)
    add_table(doc, ["Роль", "Пользователь", "Права", "Обоснование"], ROLES, widths=[3, 3.5, 6, 5])

    doc.add_heading("2.5. Реализация проекта БД", level=1)
    doc.add_heading("2.5.1. Создание таблиц", level=2)
    add_para(doc, "Таблицы, первичные ключи, внешние ключи, CHECK-ограничения и справочники реализованы в car_rental.sql. Заполнение данными выполнено только для справочников.")
    doc.add_heading("2.5.2. Создание представлений", level=2)
    add_table(doc, ["Представление", "Назначение"], VIEWS, widths=[5, 11])
    doc.add_heading("2.5.3. Назначение прав доступа", level=2)
    add_para(doc, "Права назначаются через CREATE ROLE и GRANT. Для обычных ролей доступ к sequence выдан только к тем sequence, которые нужны для вставки данных в разрешенные таблицы.")
    doc.add_heading("2.5.4. Создание триггеров", level=2)
    add_para(doc, "В проекте реализованы функции и триггеры: trg_rental_period_not_in_maintenance, trg_maintenance_not_overlap_rentals, trg_rental_set_actual_period, trg_car_must_be_active, trg_customer_license_valid_for_rental, trg_rental_opened_by_agent_or_manager, trg_maintenance_performed_by_mechanic. Дополнительно подготовлен smoke_tests.sql для проверки ключевых сценариев.")
    doc.add_heading("2.5.5. Создание индексов", level=2)
    add_para(doc, "Созданы индексы по всем внешним ключам, по дате платежа и составной индекс rental(car_id, status_code). Уникальные и первичные ключи отдельно не индексируются, так как PostgreSQL создает эти индексы автоматически.")
    doc.add_heading("2.5.6. Разработка стратегии резервного копирования", level=2)
    add_para(doc, "Полная резервная копия выполняется ежедневно ночью. При активной эксплуатации включается WAL-архивация. Хранить 14 ежедневных и 3 еженедельные полные копии, всего 17 точек восстановления, а журналы WAL хранить на тот же период.")

    doc.add_heading("Приложение А. Проверка работоспособности", level=1)
    add_para(doc, "Порядок проверки: создать пустую базу PostgreSQL; выполнить car_rental.sql; выполнить smoke_tests.sql; проверить, что ожидаемые ошибочные вставки отклоняются, а представления возвращают результат.")
    add_para(doc, "Команды: psql -d <database> -f car_rental.sql; psql -d <database> -f smoke_tests.sql")
    doc.add_heading("KPI готовности v2", level=1)
    add_table(
        doc,
        ["Критерий", "Вес", "Балл v2", "Комментарий"],
        KPI_ROWS + [("Итого", 100, sum(r[2] for r in KPI_ROWS), "Студенческий уровень близко к отличному; финальный прогон в PostgreSQL нужен для закрытия технического риска.")],
        widths=[5, 2, 2, 8],
    )
    add_para(doc, "Приложение к отчету: car_rental.sql и smoke_tests.sql находятся в папке Homework_v2.")

    docx_path = BASE / "ФАМИЛИЯ1_ФАМИЛИЯ2_ГРУППА_2.docx"
    doc.save(docx_path)
    shutil.copyfile(docx_path, BASE / "report_v2.docx")
    return docx_path


if __name__ == "__main__":
    diagram = make_er_diagram()
    write_markdown(diagram)
    write_kpi()
    result = write_docx(diagram)
    print(result)
